# ============================================================================
# comparador_objetos.py — Comparación de objetos complejos de Excel
# ============================================================================
"""
Módulo encargado de comparar objetos complejos entre la plantilla y
los libros de estudiantes:
  - Tablas de Excel (ListObjects)
  - Celdas combinadas (merged cells)
  - Validaciones de datos
  - Reglas de formato condicional
  - Dimensiones de filas y columnas
  - Gráficos (Charts) — detección básica con openpyxl
"""

from openpyxl.comments import Comment
from openpyxl.styles import PatternFill
from config import COLOR_ERROR_FILL, MENSAJES

RELLENO_ERROR = PatternFill(
    start_color=COLOR_ERROR_FILL,
    end_color=COLOR_ERROR_FILL,
    fill_type="solid"
)


# ============================================================================
# TABLAS DE EXCEL
# ============================================================================
def comparar_tablas(ws_plantilla, ws_estudiante):
    """
    Compara las tablas de Excel (Table objects) entre plantilla y estudiante.

    Retorna:
        tuple: (errores_count: int, detalles: list[str])
    """
    errores = 0
    detalles = []

    tablas_plantilla = {t.displayName: t for t in ws_plantilla.tables.values()}
    tablas_estudiante = {t.displayName: t for t in ws_estudiante.tables.values()}

    # Verificar tablas que existen en la plantilla pero no en el estudiante
    for nombre, tabla_p in tablas_plantilla.items():
        if nombre not in tablas_estudiante:
            msg = MENSAJES["tabla_faltante"].format(nombre=nombre)
            detalles.append(msg)
            errores += 1
            # Intentar marcar la primera celda del rango de la tabla
            try:
                rango = tabla_p.ref
                primera_celda = rango.split(":")[0]
                celda = ws_estudiante[primera_celda]
                celda.fill = RELLENO_ERROR
                celda.comment = Comment(text=msg, author="Auditor Excel")
            except Exception:
                pass
        else:
            tabla_e = tablas_estudiante[nombre]

            # Comparar rango de la tabla
            if tabla_p.ref != tabla_e.ref:
                msg = MENSAJES["tabla_rango"].format(
                    esperado=tabla_p.ref,
                    encontrado=tabla_e.ref
                )
                detalles.append(f"Tabla '{nombre}': {msg}")
                errores += 1

            # Comparar estilo de la tabla
            estilo_p = tabla_p.tableStyleInfo
            estilo_e = tabla_e.tableStyleInfo
            if estilo_p and estilo_e:
                if estilo_p.name != estilo_e.name:
                    msg = MENSAJES["tabla_estilo"].format(
                        esperado=estilo_p.name or "ninguno",
                        encontrado=estilo_e.name or "ninguno"
                    )
                    detalles.append(f"Tabla '{nombre}': {msg}")
                    errores += 1
            elif estilo_p and not estilo_e:
                detalles.append(f"Tabla '{nombre}': ❌ Falta estilo de tabla.")
                errores += 1

    return errores, detalles


# ============================================================================
# CELDAS COMBINADAS (MERGED CELLS)
# ============================================================================
def comparar_celdas_combinadas(ws_plantilla, ws_estudiante):
    """
    Compara los rangos de celdas combinadas (merged cells).

    Retorna:
        tuple: (errores_count: int, detalles: list[str])
    """
    errores = 0
    detalles = []

    merges_p = set(str(m) for m in ws_plantilla.merged_cells.ranges)
    merges_e = set(str(m) for m in ws_estudiante.merged_cells.ranges)

    # Rangos combinados que faltan en el estudiante
    faltantes = merges_p - merges_e
    for rango in faltantes:
        msg = MENSAJES["merge_faltante"].format(rango=rango)
        detalles.append(msg)
        errores += 1
        # Intentar marcar la primera celda del rango
        try:
            primera_celda = rango.split(":")[0]
            celda = ws_estudiante[primera_celda]
            celda.fill = RELLENO_ERROR
            celda.comment = Comment(text=msg, author="Auditor Excel")
        except Exception:
            pass

    # Rangos combinados extra en el estudiante
    extras = merges_e - merges_p
    for rango in extras:
        msg = MENSAJES["merge_extra"].format(rango=rango)
        detalles.append(msg)
        errores += 1

    return errores, detalles


# ============================================================================
# VALIDACIONES DE DATOS
# ============================================================================
def comparar_validaciones(ws_plantilla, ws_estudiante):
    """
    Compara las reglas de validación de datos entre plantilla y estudiante.

    Retorna:
        tuple: (errores_count: int, detalles: list[str])
    """
    errores = 0
    detalles = []

    # Obtener validaciones de la plantilla
    validaciones_p = {}
    if ws_plantilla.data_validations:
        for dv in ws_plantilla.data_validations.dataValidation:
            # Usar la representación del rango como clave
            rango_str = str(dv.sqref)
            validaciones_p[rango_str] = dv

    # Obtener validaciones del estudiante
    validaciones_e = {}
    if ws_estudiante.data_validations:
        for dv in ws_estudiante.data_validations.dataValidation:
            rango_str = str(dv.sqref)
            validaciones_e[rango_str] = dv

    # Verificar validaciones que faltan en el estudiante
    for rango, dv_p in validaciones_p.items():
        if rango not in validaciones_e:
            msg = MENSAJES["validacion_faltante"].format(rango=rango)
            detalles.append(msg)
            errores += 1
            # Marcar la primera celda del rango
            try:
                primera_celda = rango.split(":")[0].split()[0]
                celda = ws_estudiante[primera_celda]
                celda.fill = RELLENO_ERROR
                celda.comment = Comment(text=msg, author="Auditor Excel")
            except Exception:
                pass
        else:
            dv_e = validaciones_e[rango]
            # Comparar configuración de la validación
            diferencias = _comparar_config_validacion(dv_p, dv_e)
            if diferencias:
                msg = MENSAJES["validacion_config"].format(
                    esperado=diferencias["esperado"],
                    encontrado=diferencias["encontrado"]
                )
                detalles.append(f"Validación en '{rango}': {msg}")
                errores += 1

    return errores, detalles


def _comparar_config_validacion(dv_p, dv_e):
    """Compara la configuración detallada de dos reglas de validación."""
    diffs = []

    if dv_p.type != dv_e.type:
        diffs.append(f"Tipo: {dv_p.type} vs {dv_e.type}")
    if str(dv_p.formula1) != str(dv_e.formula1):
        diffs.append(f"Fórmula1: {dv_p.formula1} vs {dv_e.formula1}")
    if str(dv_p.formula2) != str(dv_e.formula2):
        diffs.append(f"Fórmula2: {dv_p.formula2} vs {dv_e.formula2}")
    if dv_p.operator != dv_e.operator:
        diffs.append(f"Operador: {dv_p.operator} vs {dv_e.operator}")
    if dv_p.allow_blank != dv_e.allow_blank:
        diffs.append(f"Permitir vacío: {dv_p.allow_blank} vs {dv_e.allow_blank}")

    if diffs:
        return {
            "esperado": " | ".join(d.split(" vs ")[0] for d in diffs),
            "encontrado": " | ".join(d.split(" vs ")[1] for d in diffs)
        }
    return None


# ============================================================================
# FORMATO CONDICIONAL
# ============================================================================
def comparar_formato_condicional(ws_plantilla, ws_estudiante):
    """
    Compara las reglas de formato condicional entre plantilla y estudiante.

    Retorna:
        tuple: (errores_count: int, detalles: list[str])
    """
    errores = 0
    detalles = []

    # Recopilar reglas de la plantilla indexadas por rango
    # La API de openpyxl expone conditional_formatting como una lista de
    # objetos ConditionalFormatting, cada uno con .cells (rango) y .rules (lista)
    reglas_p = {}
    for cf in ws_plantilla.conditional_formatting:
        rango_str = str(cf.cells) if hasattr(cf, 'cells') else str(cf.sqref) if hasattr(cf, 'sqref') else str(cf)
        reglas_list = cf.rules if hasattr(cf, 'rules') else cf.cfRule if hasattr(cf, 'cfRule') else []
        reglas_p[rango_str] = list(reglas_list)

    # Recopilar reglas del estudiante
    reglas_e = {}
    for cf in ws_estudiante.conditional_formatting:
        rango_str = str(cf.cells) if hasattr(cf, 'cells') else str(cf.sqref) if hasattr(cf, 'sqref') else str(cf)
        reglas_list = cf.rules if hasattr(cf, 'rules') else cf.cfRule if hasattr(cf, 'cfRule') else []
        reglas_e[rango_str] = list(reglas_list)

    # Verificar reglas faltantes
    for rango_str in reglas_p:
        if rango_str not in reglas_e:
            msg = MENSAJES["formato_cond"].format(rango=rango_str)
            detalles.append(msg)
            errores += 1
            # Marcar primera celda
            try:
                primera_celda = rango_str.split(":")[0].split()[0]
                celda = ws_estudiante[primera_celda]
                celda.fill = RELLENO_ERROR
                celda.comment = Comment(text=msg, author="Auditor Excel")
            except Exception:
                pass
        else:
            # Comparar cantidad de reglas y tipos
            reglas_pl = reglas_p[rango_str]
            reglas_el = reglas_e[rango_str]

            # Comparar tipos de reglas
            tipos_p = sorted([r.type for r in reglas_pl if hasattr(r, 'type')])
            tipos_e = sorted([r.type for r in reglas_el if hasattr(r, 'type')])

            if tipos_p != tipos_e:
                msg = (f"❌ Formato condicional en '{rango_str}': "
                       f"Tipos esperados: {tipos_p} | Encontrados: {tipos_e}")
                detalles.append(msg)
                errores += 1

    return errores, detalles


# ============================================================================
# DIMENSIONES DE FILAS Y COLUMNAS
# ============================================================================
def comparar_dimensiones(ws_plantilla, ws_estudiante, tolerancia=0.5):
    """
    Compara anchos de columna y altos de fila entre plantilla y estudiante.

    Parámetros:
        tolerancia: Margen aceptable de diferencia en las dimensiones.

    Retorna:
        tuple: (errores_count: int, detalles: list[str])
    """
    errores = 0
    detalles = []

    # Comparar anchos de columna
    for col_letter, dim_p in ws_plantilla.column_dimensions.items():
        dim_e = ws_estudiante.column_dimensions.get(col_letter)
        ancho_p = dim_p.width if dim_p.width is not None else 8.43  # Ancho predeterminado
        ancho_e = dim_e.width if (dim_e and dim_e.width is not None) else 8.43

        if abs(ancho_p - ancho_e) > tolerancia:
            msg = MENSAJES["ancho_columna"].format(
                esperado=f"{ancho_p:.2f}",
                encontrado=f"{ancho_e:.2f}"
            )
            detalles.append(f"Columna '{col_letter}': {msg}")
            errores += 1

    # Comparar altos de fila
    for row_num, dim_p in ws_plantilla.row_dimensions.items():
        dim_e = ws_estudiante.row_dimensions.get(row_num)
        alto_p = dim_p.height if dim_p.height is not None else 15.0  # Alto predeterminado
        alto_e = dim_e.height if (dim_e and dim_e.height is not None) else 15.0

        if abs(alto_p - alto_e) > tolerancia:
            msg = MENSAJES["alto_fila"].format(
                esperado=f"{alto_p:.2f}",
                encontrado=f"{alto_e:.2f}"
            )
            detalles.append(f"Fila {row_num}: {msg}")
            errores += 1

    return errores, detalles


# ============================================================================
# GRÁFICOS (Charts) — Comparación básica con openpyxl
# ============================================================================
def comparar_graficos_openpyxl(ws_plantilla, ws_estudiante):
    """
    Compara la presencia y propiedades básicas de gráficos usando openpyxl.
    Nota: openpyxl tiene soporte limitado para gráficos; para inspección
    profunda se usa win32com en el módulo comparador_com.py.

    Retorna:
        tuple: (errores_count: int, detalles: list[str])
    """
    errores = 0
    detalles = []

    charts_p = ws_plantilla._charts if hasattr(ws_plantilla, '_charts') else []
    charts_e = ws_estudiante._charts if hasattr(ws_estudiante, '_charts') else []

    cant_p = len(charts_p)
    cant_e = len(charts_e)

    if cant_p != cant_e:
        msg = (f"❌ Cantidad de gráficos incorrecta. "
               f"Esperados: {cant_p} | Encontrados: {cant_e}")
        detalles.append(msg)
        errores += abs(cant_p - cant_e)

    # Comparar gráficos que sí existen en ambos (por posición)
    for i in range(min(cant_p, cant_e)):
        chart_p = charts_p[i]
        chart_e = charts_e[i]

        # Tipo de gráfico
        tipo_p = type(chart_p).__name__
        tipo_e = type(chart_e).__name__
        if tipo_p != tipo_e:
            msg = (f"❌ Gráfico #{i+1}: Tipo incorrecto. "
                   f"Esperado: {tipo_p} | Encontrado: {tipo_e}")
            detalles.append(msg)
            errores += 1

        # Título del gráfico
        titulo_p = str(chart_p.title) if chart_p.title else "(sin título)"
        titulo_e = str(chart_e.title) if chart_e.title else "(sin título)"
        if titulo_p != titulo_e:
            msg = (f"❌ Gráfico #{i+1}: Título incorrecto. "
                   f"Esperado: '{titulo_p}' | Encontrado: '{titulo_e}'")
            detalles.append(msg)
            errores += 1

        # Cantidad de series de datos
        series_p = len(chart_p.series) if hasattr(chart_p, 'series') else 0
        series_e = len(chart_e.series) if hasattr(chart_e, 'series') else 0
        if series_p != series_e:
            msg = (f"❌ Gráfico #{i+1}: Series de datos incorrectas. "
                   f"Esperadas: {series_p} | Encontradas: {series_e}")
            detalles.append(msg)
            errores += 1

    return errores, detalles
